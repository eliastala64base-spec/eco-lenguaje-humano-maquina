import hashlib
import importlib.util
import json
import shutil
import sqlite3
import subprocess
import tempfile
import unittest
import wave
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from openpyxl import Workbook
from PIL import Image
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("lateral", ROOT / "scripts" / "procesar_recurso_lateral.py")
lateral = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(lateral)


def records(path: Path):
    return lateral.read_jsonl(path.with_name(path.name + ".jsonl"))


def by_type(path: Path, record_type: str):
    return [item for item in records(path) if item["record_type"] == record_type]


def write_pdf(path: Path, text: str | None = None):
    pdf = canvas.Canvas(str(path))
    if text:
        pdf.drawString(72, 720, text)
    pdf.showPage()
    pdf.save()


def write_xlsx(path: Path, macro=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "DATOS"
    sheet["B4"] = 3.2
    sheet["F18"] = "=B4*2"
    book.save(path)
    if macro:
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"macro-no-ejecutada")


class RecursoLateralTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name) / "02_PROYECTOS_INDEPENDIENTES" / "recursos"
        self.root.mkdir(parents=True)

    def tearDown(self):
        self.temp.cleanup()

    def test_pdf_textual_creates_only_expected_bundle(self):
        source = self.root / "Manual.pdf"
        write_pdf(source, "Diseño de mezcla de concreto")
        result = lateral.process(source)
        self.assertEqual(result["estado"], "COMPLETADO")
        self.assertEqual(
            sorted(path.name for path in self.root.iterdir()),
            ["Manual.pdf", "Manual.pdf.jsonl", "Manual.pdf.md", "Manual.pdf.sha256"],
        )
        self.assertIn("Diseño de mezcla", by_type(source, "page")[0]["text"])

    def test_scanned_pdf_records_explicit_limitation(self):
        source = self.root / "Escaneado.pdf"
        write_pdf(source)
        lateral.process(source)
        self.assertIn("NO_EXTRAIDO", [item["code"] for item in by_type(source, "warning")])

    def test_docx_table_is_structured(self):
        source = self.root / "Tabla.docx"
        document = Document()
        document.add_heading("Encabezado", level=1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        document.save(source)
        lateral.process(source)
        self.assertTrue(by_type(source, "heading"))
        self.assertEqual(by_type(source, "table")[0]["rows"], [["A", "B"]])

    def test_xlsx_formulas_and_sheets_share_one_jsonl(self):
        source = self.root / "Escalera.xlsx"
        write_xlsx(source)
        lateral.process(source)
        self.assertEqual(by_type(source, "sheet")[0]["name"], "DATOS")
        self.assertEqual(by_type(source, "formula")[0]["address"], "F18")
        self.assertFalse((self.root / "Escalera.xlsx.sheet-DATOS.csv").exists())

    def test_xlsm_macro_is_detected_not_executed(self):
        source = self.root / "Macro.xlsm"
        write_xlsx(source, macro=True)
        lateral.process(source)
        self.assertIn("MACRO_NO_EJECUTADA", [item["code"] for item in by_type(source, "warning")])
        self.assertFalse(by_type(source, "security")[0]["macros_executed"])

    def test_dwg_rvt_and_etabs_never_invent_geometry(self):
        for name in ("Plano.dwg", "Modelo.rvt", "ETABS.edb"):
            source = self.root / name
            source.write_bytes(b"formato-propietario")
            lateral.process(source)
            codes = [item["code"] for item in by_type(source, "warning")]
            self.assertIn("GEOMETRIA_NO_REPRESENTADA", codes)
            self.assertFalse(by_type(source, "entity"))

    def test_image_metadata_and_limitation(self):
        source = self.root / "Foto.jpg"
        exif = Image.Exif()
        exif[274] = 6
        Image.new("RGB", (40, 30), color="red").save(source, exif=exif)
        lateral.process(source)
        image = by_type(source, "image")[0]
        self.assertEqual((image["width"], image["height"]), (40, 30))
        self.assertEqual(image["orientation"], 6)
        self.assertIn("NO_EXTRAIDO", [item["code"] for item in by_type(source, "warning")])

    def test_audio_static_probe_does_not_execute_source(self):
        source = self.root / "Audio.wav"
        with wave.open(str(source), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(8000)
            audio.writeframes(b"\x00\x00" * 800)
        lateral.process(source)
        self.assertTrue(by_type(source, "audio_segment"))
        self.assertFalse(by_type(source, "security")[0]["source_executed"])

    def test_video_failure_is_a_warning_not_an_invention(self):
        source = self.root / "Video.mp4"
        source.write_bytes(b"\x00\x00\x00\x18ftypmp42dummy")
        lateral.process(source)
        self.assertIn("EXTRACCION_PARCIAL", [item["code"] for item in by_type(source, "warning")])

    def test_zip_inventory_without_extraction_or_markdown(self):
        source = self.root / "Paquete.zip"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("docs/a.txt", "contenido")
        lateral.process(source)
        self.assertEqual(by_type(source, "archive_member")[0]["path"], "docs/a.txt")
        self.assertFalse(source.with_name(source.name + ".md").exists())

    def test_eml_extracts_message_and_attachment(self):
        source = self.root / "Correo.eml"
        source.write_bytes(
            b"From: a@example.com\nTo: b@example.com\nSubject: Prueba\nMIME-Version: 1.0\n"
            b"Content-Type: multipart/mixed; boundary=x\n\n--x\nContent-Type: text/plain\n\nCuerpo\n"
            b"--x\nContent-Type: text/plain\nContent-Disposition: attachment; filename=a.txt\n\nAdjunto\n--x--\n"
        )
        lateral.process(source)
        self.assertEqual(by_type(source, "email_message")[0]["subject"], "Prueba")
        self.assertEqual(by_type(source, "attachment")[0]["filename"], "a.txt")

    def test_sqlite_is_opened_read_only_and_schema_is_extracted(self):
        source = self.root / "Datos.sqlite"
        connection = sqlite3.connect(source)
        connection.execute("CREATE TABLE items(id INTEGER PRIMARY KEY, name TEXT)")
        connection.execute("INSERT INTO items(name) VALUES ('uno')")
        connection.commit()
        connection.close()
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        lateral.process(source)
        self.assertEqual(by_type(source, "database_table")[0]["row_count"], 1)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_script_is_parsed_statically_and_secret_value_not_in_markdown(self):
        source = self.root / "programa.py"
        source.write_text("import os\nAPI_KEY='valor-secreto'\ndef main():\n    return 1\n", encoding="utf-8")
        lateral.process(source)
        self.assertEqual(by_type(source, "code_symbol")[0]["name"], "main")
        self.assertEqual(by_type(source, "dependency")[0]["name"], "os")
        markdown = source.with_name(source.name + ".md").read_text(encoding="utf-8")
        self.assertNotIn("valor-secreto", markdown)

    def test_executable_gets_jsonl_and_sha_only(self):
        source = self.root / "instalador.exe"
        source.write_bytes(b"MZ" + b"\x00" * 64)
        lateral.process(source)
        self.assertTrue(source.with_name(source.name + ".jsonl").is_file())
        self.assertTrue(source.with_name(source.name + ".sha256").is_file())
        self.assertFalse(source.with_name(source.name + ".md").exists())
        self.assertEqual(by_type(source, "file_role")[0]["role"], "EJECUTABLE_ACTIVO")

    def test_exact_duplicate_and_renamed_copy_get_sha_only_for_second(self):
        first = self.root / "A.pdf"
        second = self.root / "Copia renombrada.pdf"
        write_pdf(first, "igual")
        second.write_bytes(first.read_bytes())
        lateral.process(self.root)
        self.assertTrue(first.with_name(first.name + ".jsonl").exists())
        self.assertTrue(second.with_name(second.name + ".sha256").exists())
        self.assertFalse(second.with_name(second.name + ".jsonl").exists())

    def test_manifestation_requires_exact_basename_and_content_evidence(self):
        first = self.root / "MISMO.txt"
        second = self.root / "MISMO.md"
        first.write_text("contenido equivalente", encoding="utf-8")
        second.write_text("contenido   equivalente\n", encoding="utf-8")
        lateral.process(self.root)
        roles = {source.name: by_type(source, "file_role")[0]["role"] for source in (first, second)}
        self.assertEqual(sorted(roles.values()), ["FUENTE_PRIMARIA", "MANIFESTACION_ASOCIADA"])
        associated = next(source for source in (first, second) if roles[source.name] == "MANIFESTACION_ASOCIADA")
        self.assertEqual(by_type(associated, "relation")[0]["status"], "CONFIRMADA")

    def test_different_basename_never_becomes_manifestation(self):
        first = self.root / "PLANO.txt"
        second = self.root / "PLANO_FINAL.md"
        first.write_text("contenido equivalente", encoding="utf-8")
        second.write_text("contenido   equivalente\n", encoding="utf-8")
        lateral.process(self.root)
        self.assertEqual(by_type(first, "file_role")[0]["role"], "FUENTE_PRIMARIA")
        self.assertEqual(by_type(second, "file_role")[0]["role"], "FUENTE_PRIMARIA")

    def test_auxiliary_log_has_jsonl_sha_and_useful_markdown(self):
        source = self.root / "MODELO.log"
        source.write_text("Análisis terminado sin errores", encoding="utf-8")
        lateral.process(source)
        self.assertEqual(by_type(source, "file_role")[0]["role"], "ARCHIVO_AUXILIAR")
        self.assertTrue(source.with_name(source.name + ".md").exists())

    def test_temporary_has_no_derivatives(self):
        source = self.root / "~$bloqueo.tmp"
        source.write_bytes(b"temporal")
        result = lateral.process(source)
        self.assertEqual(result["resultados"][0]["status"], "DETECTADO_SIN_DERIVADOS")
        self.assertEqual([path.name for path in self.root.iterdir()], [source.name])

    def test_technical_family_assigns_primary_generated_auxiliary_and_backup(self):
        edb = self.root / "MODELO.edb"
        xlsx = self.root / "MODELO.xlsx"
        log = self.root / "MODELO.log"
        backup = self.root / "MODELO.bak"
        edb.write_bytes(b"edb")
        write_xlsx(xlsx)
        log.write_text("log", encoding="utf-8")
        backup.write_bytes(b"backup")
        lateral.process(self.root)
        self.assertEqual(by_type(edb, "file_role")[0]["role"], "FUENTE_PRIMARIA")
        self.assertEqual(by_type(xlsx, "file_role")[0]["role"], "DERIVADO_GENERADO")
        self.assertEqual(by_type(log, "file_role")[0]["role"], "ARCHIVO_AUXILIAR")
        self.assertFalse(backup.with_name(backup.name + ".jsonl").exists())
        self.assertTrue(backup.with_name(backup.name + ".sha256").exists())

    def test_existing_80_is_skipped_and_never_processed(self):
        output = self.root / lateral.OUTPUT_80
        output.mkdir()
        (output / "derivado.pdf").write_bytes(b"no procesar")
        source = self.root / "Libro.pdf"
        write_pdf(source, "libro")
        lateral.process(self.root)
        self.assertFalse((output / "derivado.pdf.jsonl").exists())
        self.assertFalse((self.root / lateral.OUTPUT_80 / lateral.OUTPUT_80).exists())

    def test_orphaned_self_derivatives_are_not_processed_recursively(self):
        source = self.root / "Libro.txt"
        source.write_text("contenido", encoding="utf-8")
        lateral.process(source)
        source.unlink()
        lateral.process(self.root)
        names = {path.name for path in self.root.iterdir()}
        self.assertFalse(any(name.endswith(".jsonl.jsonl") or name.endswith(".md.jsonl") for name in names))

    def test_execution_inside_80_is_blocked(self):
        output = self.root / lateral.OUTPUT_80
        output.mkdir()
        result = lateral.process(output)
        self.assertEqual(result["estado"], "BLOQUEADO")

    def test_json_source_uses_full_original_name_as_prefix(self):
        source = self.root / "configuracion.json"
        source.write_text('{"activo":true}', encoding="utf-8")
        lateral.process(source)
        self.assertTrue((self.root / "configuracion.json.jsonl").exists())
        self.assertTrue((self.root / "configuracion.json.md").exists())
        self.assertTrue((self.root / "configuracion.json.sha256").exists())

    def test_jsonl_hierarchy_and_universal_records(self):
        source = self.root / "texto.txt"
        source.write_text("material técnico", encoding="utf-8")
        lateral.process(source)
        result = records(source)
        self.assertTrue(lateral.UNIVERSAL.issubset({item["record_type"] for item in result}))
        ranks = [lateral.ORDER.get(item["record_type"], 7) for item in result]
        self.assertEqual(ranks, sorted(ranks))

    def test_universal_machine_paths_are_identical_across_formats(self):
        pdf = self.root / "Manual.pdf"
        spreadsheet = self.root / "Calculo.xlsx"
        write_pdf(pdf, "Diseño de mezcla")
        write_xlsx(spreadsheet)
        for source in (pdf, spreadsheet):
            lateral.process(source)
            source_metadata = by_type(source, "source_metadata")[0]
            integrity = by_type(source, "integrity")[0]
            summary = by_type(source, "summary")[0]
            keywords = by_type(source, "keywords")[0]
            self.assertEqual(source_metadata["source_size_bytes"], source.stat().st_size)
            self.assertEqual(integrity["source_size_bytes"], source_metadata["source_size_bytes"])
            self.assertEqual(integrity["source_sha256"], lateral.sha256_file(source))
            self.assertEqual(summary["text"], summary["brief_description"])
            for field in ("short_description", "brief_description", "long_description"):
                self.assertTrue(summary[field])
            self.assertIsInstance(summary["synonyms"], list)
            self.assertIsInstance(keywords["values"], list)

    def test_markdown_and_sha256_repeat_stable_integrity_tags(self):
        source = self.root / "Etiquetas.txt"
        source.write_text("contenido técnico", encoding="utf-8")
        lateral.process(source)
        digest = lateral.sha256_file(source)
        markdown = source.with_name(source.name + ".md").read_text(encoding="utf-8")
        checksum_file = source.with_name(source.name + ".sha256").read_text(encoding="utf-8")
        self.assertIn(f"schema_version: {lateral.SCHEMA_VERSION}\n", markdown)
        self.assertIn(f"source_size_bytes: {source.stat().st_size}\n", markdown)
        self.assertIn(f"source_sha256: {digest}\n", markdown)
        self.assertIn(f"# schema_name: {lateral.SCHEMA_NAME}\n", checksum_file)
        self.assertIn("# artifact_type: integrity\n", checksum_file)
        self.assertIn(f"# source_size_bytes: {source.stat().st_size}\n", checksum_file)
        self.assertIn(f"# source_sha256: {digest}\n", checksum_file)
        self.assertTrue(checksum_file.endswith(f"{digest}  {source.name}\n"))
        if shutil.which("sha256sum"):
            checked = subprocess.run(
                ["sha256sum", "-c", source.name + ".sha256"],
                cwd=self.root,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(checked.returncode, 0, checked.stderr)

    def test_validator_rejects_inconsistent_universal_size(self):
        source = self.root / "Tamano.txt"
        source.write_text("12345", encoding="utf-8")
        lateral.process(source)
        jsonl = source.with_name(source.name + ".jsonl")
        result = lateral.read_jsonl(jsonl)
        next(item for item in result if item["record_type"] == "source_metadata")["source_size_bytes"] += 1
        lateral.write_jsonl(jsonl, result)
        with self.assertRaisesRegex(ValueError, "Metadatos universales"):
            lateral.validate_jsonl(jsonl, source)

    def test_markdown_is_regenerated_from_validated_jsonl(self):
        source = self.root / "texto.md"
        source.write_text("consulta Obsidian", encoding="utf-8")
        lateral.process(source)
        jsonl = source.with_name(source.name + ".jsonl")
        expected = lateral.render_markdown(lateral.validate_jsonl(jsonl, source), lateral.sha256_file(jsonl))
        self.assertEqual(source.with_name(source.name + ".md").read_text(encoding="utf-8"), expected)

    def test_single_ai_campaign_enrichment_requires_evidence(self):
        source = self.root / "Norma.txt"
        source.write_text("Norma técnica para concreto", encoding="utf-8")
        enrichment = self.root.parent / "enrichment.json"
        enrichment.write_text(json.dumps({
            "Norma.txt": {
                "short_description": "Norma de concreto.",
                "brief_description": "Norma técnica sobre concreto.",
                "long_description": "Documento normativo que presenta requisitos técnicos observados para concreto.",
                "synonyms": ["reglamento de concreto"],
                "keywords": ["concreto"],
                "entities": [{"name": "concreto", "entity_type": "MATERIAL", "evidence": "texto:1"}],
                "relations": [],
            }
        }), encoding="utf-8")
        lateral.process(source, enrichment)
        quality = by_type(source, "quality")[0]
        self.assertTrue(quality["ai_enrichment_applied"])
        self.assertEqual(quality["ai_passes"], 1)
        summary = by_type(source, "summary")[0]
        self.assertEqual(summary["generation_method"], "AI_UNICA")
        self.assertEqual(summary["short_description"], "Norma de concreto.")
        self.assertEqual(summary["brief_description"], "Norma técnica sobre concreto.")
        self.assertEqual(summary["long_description"], "Documento normativo que presenta requisitos técnicos observados para concreto.")
        self.assertEqual(summary["synonyms"], ["reglamento de concreto"])

    def test_ai_relation_without_evidence_is_rejected_without_publication(self):
        source = self.root / "Norma.txt"
        source.write_text("Norma técnica", encoding="utf-8")
        enrichment = self.root.parent / "bad-enrichment.json"
        enrichment.write_text(json.dumps({"Norma.txt": {"relations": [{"relation_type": "RELACIONADO_CON"}]}}), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "sin evidencia"):
            lateral.process(source, enrichment)
        self.assertFalse(source.with_name(source.name + ".jsonl").exists())

    def test_validation_failure_preserves_existing_derivatives_and_original(self):
        source = self.root / "Seguro.txt"
        source.write_text("original", encoding="utf-8")
        old_jsonl = source.with_name(source.name + ".jsonl")
        old_jsonl.write_text("anterior\n", encoding="utf-8")
        source_hash = lateral.sha256_file(source)
        with self.assertRaisesRegex(ValueError, "Fallo de validación"):
            lateral.process(source, forced_failure=True)
        self.assertEqual(old_jsonl.read_text(encoding="utf-8"), "anterior\n")
        self.assertEqual(lateral.sha256_file(source), source_hash)

    def test_publication_failure_restores_all_previous_lateral_derivatives(self):
        source = self.root / "Seguro.txt"
        source.write_text("original", encoding="utf-8")
        destinations = lateral.output_paths(source)
        destinations["jsonl"].write_text("jsonl anterior", encoding="utf-8")
        destinations["markdown"].write_text("markdown anterior", encoding="utf-8")
        destinations["sha256"].write_text("sha anterior", encoding="utf-8")
        stage = Path(tempfile.mkdtemp(dir=self.root))
        for key, destination in destinations.items():
            (stage / destination.name).write_text(f"nuevo {key}", encoding="utf-8")
        real_replace = lateral.os.replace

        def fail_first_publication(source_path, destination_path):
            if Path(source_path) == stage / destinations["jsonl"].name:
                raise OSError("fallo simulado")
            return real_replace(source_path, destination_path)

        with mock.patch.object(lateral.os, "replace", side_effect=fail_first_publication):
            with self.assertRaisesRegex(OSError, "fallo simulado"):
                lateral.publish_bundle(stage, source, {"jsonl", "markdown", "sha256"})
        self.assertEqual(destinations["jsonl"].read_text(encoding="utf-8"), "jsonl anterior")
        self.assertEqual(destinations["markdown"].read_text(encoding="utf-8"), "markdown anterior")
        self.assertEqual(destinations["sha256"].read_text(encoding="utf-8"), "sha anterior")

    def test_no_txt_csv_sqlite_or_separate_json_derivatives_are_created(self):
        source = self.root / "datos.csv"
        source.write_text("a,b\n1,2\n", encoding="utf-8")
        lateral.process(source)
        names = {path.name for path in self.root.iterdir()}
        self.assertNotIn("datos.csv.txt", names)
        self.assertNotIn("datos.csv.structure.json", names)
        self.assertNotIn("datos.csv.sqlite", names)
        self.assertEqual([name for name in names if name.endswith(".jsonl")], ["datos.csv.jsonl"])


if __name__ == "__main__":
    unittest.main()
