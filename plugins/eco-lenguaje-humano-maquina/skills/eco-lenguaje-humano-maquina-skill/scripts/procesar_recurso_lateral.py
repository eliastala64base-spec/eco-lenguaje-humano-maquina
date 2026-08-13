#!/usr/bin/env python3
"""Procesa recursos externos en modo lateral sin ejecutar contenido activo."""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import subprocess
import tarfile
import tempfile
import wave
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import quote
from xml.etree import ElementTree as ET

SKILL_NAME = "Lenguaje Humano–Máquina"
SKILL_VERSION = "0.5.0-beta.1"
SCHEMA_NAME = "lenguaje_humano_maquina"
SCHEMA_VERSION = 2
PROCESSING_MODE = "RECURSO_PORTABLE_LATERAL"
OUTPUT_80 = "80_LENGUAJE_HUMANO_MAQUINA"
MAX_TEXT_CHARS = 5_000_000
MAX_ARCHIVE_MEMBERS = 100_000
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")

DOCUMENTS = {".pdf", ".docx", ".doc", ".docm", ".rtf", ".odt", ".txt", ".md", ".epub"}
SPREADSHEETS = {".xlsx", ".xls", ".xlsm", ".xlsb", ".csv", ".ods"}
PRESENTATIONS = {".pptx", ".ppt", ".pptm", ".odp"}
IMAGES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp", ".heic", ".gif"}
AUDIO = {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg", ".wma"}
VIDEO = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".mpeg", ".mpg", ".m4v"}
TECHNICAL = {
    ".dwg", ".dxf", ".dwt", ".dwf", ".dwfx", ".dgn", ".rvt", ".rfa", ".ifc",
    ".nwd", ".nwc", ".edb", ".sdb", ".f2k", ".e2k", ".qgz", ".qgs", ".shp",
    ".geojson", ".kml", ".kmz", ".las", ".laz", ".stl", ".obj", ".fbx", ".3ds",
}
ARCHIVES = {".zip", ".rar", ".7z", ".tar", ".gz", ".tgz", ".bz2", ".xz", ".cab", ".dmg", ".img", ".kmz"}
EMAILS = {".eml", ".msg"}
DATABASES = {".sqlite", ".sqlite3", ".db", ".accdb", ".mdb"}
CODE = {
    ".py", ".pyi", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".java", ".c",
    ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt",
    ".kts", ".scala", ".sql", ".r", ".m", ".vue", ".svelte", ".html", ".css",
    ".scss", ".json", ".json5", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".xml",
}
ACTIVE = {".exe", ".dll", ".msi", ".sys", ".bat", ".cmd", ".ps1", ".sh", ".vbs", ".jar", ".wasm", ".com", ".scr"}
BACKUPS = {".bak", ".backup", ".old", ".orig", ".bk"}
AUXILIARY = {".log", ".xsd", ".properties", ".prj", ".cpg", ".dbf", ".shx", ".sld", ".aux"}
TEMPORARY = {".tmp", ".temp", ".cache", ".lock", ".lck", ".swp", ".part", ".crdownload", ".autosave"}
ALL_KNOWN = DOCUMENTS | SPREADSHEETS | PRESENTATIONS | IMAGES | AUDIO | VIDEO | TECHNICAL | ARCHIVES | EMAILS | DATABASES | CODE | ACTIVE | BACKUPS | AUXILIARY
TECHNICAL_PRIMARY = {".edb", ".sdb", ".f2k", ".e2k", ".rvt", ".nwd", ".qgz", ".qgs", ".shp", ".dwg", ".dgn", ".ifc"}
GENERATED_CANDIDATES = {".xlsx", ".xls", ".xlsm", ".xlsb", ".csv", ".ods", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}

ORDER = {
    "manifest": 1, "source_metadata": 2, "integrity": 3, "file_role": 4,
    "format_information": 5, "structure": 6, "page": 6, "section": 6, "heading": 6,
    "table": 6, "sheet": 6, "slide": 6, "layer": 6, "archive_member": 6, "database_table": 6,
    "database_field": 6, "content": 7, "paragraph": 7, "cell": 7, "subtitle": 7,
    "transcript_segment": 7, "email_message": 7, "comment": 7, "technical_data": 8,
    "formula": 8, "chart": 8, "image": 8, "drawing": 8, "block": 8, "entity": 9,
    "bim_element": 9, "gis_feature": 9, "audio_segment": 9, "video_segment": 9,
    "attachment": 9, "database_relation": 10, "code_file": 9, "code_symbol": 9,
    "dependency": 10, "relation": 10, "summary": 11, "keywords": 12,
    "rights_observed": 13, "security": 14, "comparison_features": 15, "quality": 16,
    "warning": 17, "warnings": 17, "end": 18,
}
UNIVERSAL = {
    "manifest", "source_metadata", "integrity", "file_role", "format_information",
    "summary", "keywords", "quality", "warnings", "end",
}
REQUIRED_BY_TYPE = {
    "manifest": {"schema_name", "schema_version", "source_filename", "source_relative_path", "generated_at", "generated_by", "processing_mode"},
    "source_metadata": {"source_filename", "source_relative_path", "source_size_bytes", "source_modified_at"},
    "integrity": {"algorithm", "source_sha256", "source_size_bytes"},
    "file_role": {"role", "basis", "classification_is_final"},
    "format_information": {"declared_extension", "declared_format", "detected_format"},
    "summary": {
        "text", "short_description", "brief_description", "long_description",
        "synonyms", "generation_method",
    },
    "keywords": {"values", "classification_is_final"},
    "quality": {"extraction_confidence", "extraction_status", "deterministic_extraction", "ai_enrichment_applied", "ai_passes"},
    "warning": {"code", "description", "severity"},
    "warnings": {"codes", "count"},
    "end": {"status", "source_sha256"},
}
ALLOWED_ROLES = {
    "FUENTE_PRIMARIA", "MANIFESTACION_ASOCIADA", "ARCHIVO_AUXILIAR", "DERIVADO_GENERADO",
    "RESPALDO_COPIA", "CONTENEDOR", "BASE_DATOS", "CODIGO_FUENTE", "EJECUTABLE_ACTIVO",
    "TEMPORAL_CACHE", "DESCONOCIDO",
}
STOPWORDS = {
    "para", "como", "desde", "hasta", "entre", "sobre", "este", "esta", "estos", "estas",
    "that", "this", "with", "from", "into", "documento", "archivo", "file", "the", "and", "una",
    "uno", "unos", "unas", "del", "las", "los", "por", "con", "sin", "que", "sus", "de", "la",
    "el", "en", "un", "y", "o", "a",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, (list, tuple, set)):
        return [json_safe(item) for item in value]
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    return str(value)


def iso_timestamp(value: float | None) -> str | None:
    return datetime.fromtimestamp(value, timezone.utc).replace(microsecond=0).isoformat() if value is not None else None


def warning(code: str, description: str, severity: str = "MEDIO") -> dict[str, Any]:
    return {"record_type": "warning", "code": code, "description": description, "severity": severity}


def is_temporary(path: Path) -> bool:
    return (
        path.suffix.lower() in TEMPORARY
        or path.name.startswith("~$")
        or path.name in {".DS_Store", "Thumbs.db"}
        or path.name.endswith("~")
    )


def is_attached_derivative(path: Path) -> bool:
    if path.name.endswith(".sha256"):
        return True
    for suffix in (".jsonl", ".sha256", ".md"):
        if path.name.endswith(suffix):
            original = path.with_name(path.name[: -len(suffix)])
            if original.is_file():
                return True
    try:
        if path.name.endswith(".jsonl"):
            with path.open("r", encoding="utf-8") as stream:
                first_line = stream.readline()
            manifest = json.loads(first_line)
            return manifest.get("record_type") == "manifest" and manifest.get("schema_name") == "lenguaje_humano_maquina"
        if path.name.endswith(".md"):
            head = path.read_text(encoding="utf-8", errors="replace")[:4096]
            return (
                (
                    head.startswith(f"---\nschema_name: {SCHEMA_NAME}\n")
                    and f"\nschema_version: {SCHEMA_VERSION}\n" in head
                )
                or head.startswith("---\nschema_version: 1\n")
            ) and (
                "\njsonl_sha256: " in head
                and "\nfile_role: " in head
            )
    except (OSError, UnicodeError, json.JSONDecodeError):
        return False
    return False


def iter_sources(target: Path) -> list[Path]:
    if target.is_file():
        return [] if is_attached_derivative(target) else [target]
    sources: list[Path] = []
    for path in sorted(target.rglob("*")):
        if OUTPUT_80 in path.parts:
            continue
        if path.is_file() and not path.is_symlink() and not is_attached_derivative(path):
            sources.append(path)
    return sources


def detect_format(path: Path) -> dict[str, Any]:
    extension = path.suffix.lower()
    head = path.read_bytes()[:32]
    detected = None
    if head.startswith(b"%PDF-"):
        detected = "PDF"
    elif head.startswith(b"PK\x03\x04"):
        detected = "ZIP_OOXML_OR_CONTAINER"
    elif head.startswith(b"SQLite format 3\x00"):
        detected = "SQLITE3"
    elif head.startswith(b"MZ"):
        detected = "PE_EXECUTABLE"
    elif head.startswith(b"\x7fELF"):
        detected = "ELF_BINARY"
    elif head.startswith(b"\x89PNG\r\n\x1a\n"):
        detected = "PNG"
    elif head.startswith(b"\xff\xd8\xff"):
        detected = "JPEG"
    elif head[:4] in {b"II*\x00", b"MM\x00*"}:
        detected = "TIFF"
    elif head.startswith(b"GIF8"):
        detected = "GIF"
    elif head.startswith(b"RIFF") and head[8:12] == b"WAVE":
        detected = "WAV"
    elif head.startswith(b"RIFF") and head[8:12] == b"WEBP":
        detected = "WEBP"
    elif head.startswith(b"Rar!\x1a\x07"):
        detected = "RAR"
    elif head.startswith(b"7z\xbc\xaf\x27\x1c"):
        detected = "7Z"
    mime, _ = mimetypes.guess_type(path.name)
    return {
        "declared_extension": extension,
        "declared_format": extension[1:].upper() if extension else "SIN_EXTENSION",
        "detected_format": detected or (mime.upper() if mime else "NO_DETERMINADO"),
        "mime_type": mime,
        "signature_match": bool(detected),
    }


def base_metadata(path: Path, root: Path) -> dict[str, Any]:
    stat = path.stat()
    created = getattr(stat, "st_birthtime", None)
    return {
        "record_type": "source_metadata",
        "source_filename": path.name,
        "source_extension": path.suffix.lower(),
        "source_relative_path": path.relative_to(root).as_posix() if root.is_dir() else path.name,
        "source_size_bytes": stat.st_size,
        "source_created_at": iso_timestamp(created),
        "source_modified_at": iso_timestamp(stat.st_mtime),
        "filesystem_ctime": iso_timestamp(stat.st_ctime),
        "language": None,
        "author_detected": None,
        "organization_detected": None,
        "creator_application": None,
        "modifier_application": None,
        "metadata_confidence": "VERIFICADO_SISTEMA_ARCHIVOS",
    }


def metadata_value(metadata: dict[str, Any], *names: str) -> Any:
    folded = {str(key).casefold(): value for key, value in metadata.items()}
    return next((folded[name.casefold()] for name in names if folded.get(name.casefold()) not in (None, "")), None)


def preliminary_role(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if is_temporary(path):
        return "TEMPORAL_CACHE", "PATRON_TEMPORAL_DETERMINISTA"
    if ext in BACKUPS:
        return "RESPALDO_COPIA", "EXTENSION_DE_RESPALDO"
    if ext in ACTIVE:
        return "EJECUTABLE_ACTIVO", "EXTENSION_DE_CONTENIDO_ACTIVO"
    if ext in ARCHIVES:
        return "CONTENEDOR", "FORMATO_CONTENEDOR"
    if ext in DATABASES:
        return "BASE_DATOS", "FORMATO_BASE_DE_DATOS"
    if ext in CODE:
        return "CODIGO_FUENTE", "FORMATO_CODIGO_O_CONFIGURACION"
    if ext in AUXILIARY:
        return "ARCHIVO_AUXILIAR", "FORMATO_AUXILIAR"
    return "FUENTE_PRIMARIA", "VALOR_TECNICO_O_INTELECTUAL_INDEPENDIENTE"


def quote_sql_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def extract_pdf(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    from pypdf import PdfReader
    reader = PdfReader(path)
    records: list[dict[str, Any]] = []
    texts = []
    for index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        texts.append(text)
        records.append({"record_type": "page", "page_number": index, "text": text})
    metadata = {str(k).lstrip("/"): json_safe(v) for k, v in (reader.metadata or {}).items()}
    warnings = []
    full_text = "\n\f\n".join(texts)
    if not full_text.strip():
        warnings.append(warning("NO_EXTRAIDO", "El PDF no contiene texto extraíble; puede requerir OCR verificado."))
    metrics = {"page_count": len(reader.pages), "text_length": len(full_text)}
    return records, full_text, metrics, warnings, metadata


def extract_docx(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    from docx import Document
    document = Document(path)
    records: list[dict[str, Any]] = []
    texts = []
    heading_count = 0
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text
        if not text:
            continue
        texts.append(text)
        record_type = "heading" if paragraph.style and paragraph.style.name.startswith("Heading") else "paragraph"
        heading_count += int(record_type == "heading")
        records.append({"record_type": record_type, "paragraph_number": index, "text": text, "style": paragraph.style.name if paragraph.style else None})
    for table_index, table in enumerate(document.tables, 1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        records.append({"record_type": "table", "table_number": table_index, "rows": rows})
        texts.extend(" | ".join(row) for row in rows)
    core = document.core_properties
    metadata = {
        "title": core.title or None, "author": core.author or None, "subject": core.subject or None,
        "keywords": core.keywords or None, "comments": core.comments or None,
        "created": json_safe(core.created), "modified": json_safe(core.modified),
    }
    warnings = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/comments.xml" in names:
            warnings.append(warning("COMENTARIOS_DETECTADOS", "Se detectaron comentarios; su contexto puede requerir revisión manual.", "BAJO"))
        document_xml = archive.read("word/document.xml")
        if b"<w:ins" in document_xml or b"<w:del" in document_xml:
            warnings.append(warning("CONTROL_CAMBIOS_DETECTADO", "Se detectó control de cambios; no fue aceptado ni rechazado."))
        if any(name.startswith("word/embeddings/") for name in names):
            warnings.append(warning("OBJETO_INCRUSTADO_NO_EXTRAIDO", "Se detectaron objetos incrustados que no fueron ejecutados ni extraídos."))
        if any(name.endswith("vbaProject.bin") for name in names):
            warnings.append(warning("MACRO_NO_EJECUTADA", "Se detectó una macro que no fue ejecutada."))
    full_text = "\n".join(texts)
    metrics = {"paragraph_count": len(document.paragraphs), "heading_count": heading_count, "table_count": len(document.tables), "text_length": len(full_text)}
    return records, full_text, metrics, warnings, metadata


def extract_xlsx(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    import openpyxl
    keep_vba = path.suffix.lower() == ".xlsm"
    book = openpyxl.load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba, keep_links=True)
    try:
        values_book = openpyxl.load_workbook(path, read_only=False, data_only=True, keep_vba=False, keep_links=True)
    except Exception:
        values_book = None
    records: list[dict[str, Any]] = []
    texts = []
    formula_count = table_count = chart_count = cell_count = 0
    for order, sheet in enumerate(book.worksheets, 1):
        visibility = sheet.sheet_state.upper()
        records.append({
            "record_type": "sheet", "sheet_id": f"SHT-{order:03d}", "name": sheet.title,
            "order": order, "visibility": visibility, "used_range": sheet.calculate_dimension(),
        })
        values_sheet = values_book[sheet.title] if values_book and sheet.title in values_book.sheetnames else None
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cell_count += 1
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formula_count += 1
                    calculated = values_sheet[cell.coordinate].value if values_sheet else None
                    records.append({
                        "record_type": "formula", "sheet_id": f"SHT-{order:03d}", "address": cell.coordinate,
                        "formula": str(cell.value), "calculated_value": json_safe(calculated),
                    })
                    texts.append(f"{sheet.title}!{cell.coordinate}: {cell.value}")
                else:
                    records.append({
                        "record_type": "cell", "sheet_id": f"SHT-{order:03d}", "address": cell.coordinate,
                        "value": json_safe(cell.value), "data_type": cell.data_type,
                    })
                    texts.append(f"{sheet.title}!{cell.coordinate}: {cell.value}")
        for name, table in sheet.tables.items():
            table_count += 1
            records.append({"record_type": "table", "sheet_id": f"SHT-{order:03d}", "name": name, "range": table.ref})
        for index, chart in enumerate(sheet._charts, 1):
            chart_count += 1
            records.append({"record_type": "chart", "sheet_id": f"SHT-{order:03d}", "chart_number": index, "chart_type": type(chart).__name__})
    warnings = []
    with zipfile.ZipFile(path) as archive:
        if any(name.endswith("vbaProject.bin") for name in archive.namelist()):
            warnings.append(warning("MACRO_NO_EJECUTADA", "Se detectó contenido VBA que no fue ejecutado."))
    props = book.properties
    metadata = {
        "title": props.title, "creator": props.creator, "last_modified_by": props.lastModifiedBy,
        "created": json_safe(props.created), "modified": json_safe(props.modified),
    }
    full_text = "\n".join(texts)
    metrics = {
        "sheet_count": len(book.sheetnames), "sheet_names": list(book.sheetnames), "cell_count": cell_count,
        "formula_count": formula_count, "table_count": table_count, "chart_count": chart_count,
        "text_length": len(full_text),
    }
    book.close()
    if getattr(book, "vba_archive", None):
        book.vba_archive.close()
    if values_book:
        values_book.close()
    return records, full_text, metrics, warnings, metadata


def extract_presentation(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    from pptx import Presentation
    deck = Presentation(path)
    records: list[dict[str, Any]] = []
    texts = []
    image_count = table_count = 0
    for index, slide in enumerate(deck.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                slide_texts.append(shape.text)
            if getattr(shape, "has_table", False):
                table_count += 1
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                records.append({"record_type": "table", "slide_number": index, "rows": rows})
            if getattr(shape, "shape_type", None) == 13:
                image_count += 1
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            pass
        records.append({"record_type": "slide", "slide_number": index, "text": "\n".join(slide_texts), "notes": notes})
        texts.extend(slide_texts)
        if notes:
            texts.append(notes)
    full_text = "\n".join(texts)
    metrics = {"slide_count": len(deck.slides), "table_count": table_count, "image_count": image_count, "text_length": len(full_text)}
    return records, full_text, metrics, [], {}


def extract_image(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    from PIL import ExifTags, Image
    with Image.open(path) as image:
        exif_raw = image.getexif()
        exif = {ExifTags.TAGS.get(key, str(key)): json_safe(value) for key, value in exif_raw.items()}
        record = {
            "record_type": "image", "width": image.width, "height": image.height, "format": image.format,
            "mode": image.mode, "orientation": exif.get("Orientation"), "exif": exif,
        }
        metrics = {"image_count": 1, "width": image.width, "height": image.height, "text_length": 0}
        metadata = {"format": image.format, "mode": image.mode, "exif_present": bool(exif)}
    warnings = [warning("NO_EXTRAIDO", "No se aplicó OCR ni descripción visual determinista; requiere una única intervención de IA o revisión humana.", "BAJO")]
    if "GPSInfo" in exif:
        warnings.append(warning("POSIBLE_GEOLOCALIZACION", "Los metadatos EXIF contienen información GPS; revisar privacidad.", "ALTO"))
    return [record], "", metrics, warnings, metadata


def extract_media(path: Path, kind: str) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    command = ["ffprobe", "-v", "error", "-show_format", "-show_streams", "-print_format", "json", str(path)]
    records: list[dict[str, Any]] = []
    warnings = []
    metadata: dict[str, Any] = {}
    try:
        completed = subprocess.run(command, capture_output=True, text=True, timeout=30, check=True)
        metadata = json.loads(completed.stdout)
        for index, stream in enumerate(metadata.get("streams", [])):
            records.append({"record_type": f"{kind}_segment", "segment_number": index + 1, "stream_metadata": json_safe(stream)})
    except Exception as exc:
        if kind == "audio" and path.suffix.lower() == ".wav":
            try:
                with wave.open(str(path), "rb") as audio:
                    frames = audio.getnframes()
                    rate = audio.getframerate()
                    stream = {
                        "codec_name": "pcm",
                        "channels": audio.getnchannels(),
                        "sample_rate": rate,
                        "sample_width_bytes": audio.getsampwidth(),
                        "frame_count": frames,
                        "duration_seconds": frames / rate if rate else 0,
                    }
                records.append({"record_type": "audio_segment", "segment_number": 1, "stream_metadata": stream})
                metadata = {"streams": [stream], "format": {"duration": stream["duration_seconds"]}}
            except Exception as fallback_exc:
                warnings.append(warning("EXTRACCION_PARCIAL", f"No fue posible obtener metadatos multimedia: {type(fallback_exc).__name__}."))
        else:
            warnings.append(warning("EXTRACCION_PARCIAL", f"No fue posible obtener metadatos multimedia: {type(exc).__name__}."))
    warnings.append(warning("NO_EXTRAIDO", "No se generó transcripción; se conserva como limitación explícita.", "BAJO"))
    format_data = metadata.get("format", {})
    metrics = {
        "duration_seconds": float(format_data.get("duration", 0) or 0),
        "stream_count": len(metadata.get("streams", [])), "text_length": 0,
    }
    return records, "", metrics, warnings, metadata


def extract_archive(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    records: list[dict[str, Any]] = []
    warnings = []
    encrypted = False
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_MEMBERS:
                warnings.append(warning("EXTRACCION_PARCIAL", f"Inventario limitado a {MAX_ARCHIVE_MEMBERS} miembros."))
            for info in infos[:MAX_ARCHIVE_MEMBERS]:
                encrypted = encrypted or bool(info.flag_bits & 0x1)
                records.append({
                    "record_type": "archive_member", "path": info.filename, "size_bytes": info.file_size,
                    "compressed_size_bytes": info.compress_size, "modified_at_container": "%04d-%02d-%02dT%02d:%02d:%02d" % info.date_time,
                    "encrypted": bool(info.flag_bits & 0x1), "is_directory": info.is_dir(),
                })
    elif tarfile.is_tarfile(path):
        with tarfile.open(path, "r:*") as archive:
            members = archive.getmembers()
            for member in members[:MAX_ARCHIVE_MEMBERS]:
                records.append({
                    "record_type": "archive_member", "path": member.name, "size_bytes": member.size,
                    "modified_at_container": iso_timestamp(member.mtime), "encrypted": False,
                    "is_directory": member.isdir(),
                })
            if len(members) > MAX_ARCHIVE_MEMBERS:
                warnings.append(warning("EXTRACCION_PARCIAL", f"Inventario limitado a {MAX_ARCHIVE_MEMBERS} miembros."))
    else:
        warnings.append(warning("FORMATO_NO_SOPORTADO_COMPLETAMENTE", "Se identificó el contenedor, pero no existe un lector estático disponible."))
    if encrypted:
        warnings.append(warning("ARCHIVO_PROTEGIDO", "El contenedor incluye miembros cifrados; no fueron extraídos."))
    metrics = {"archive_member_count": len(records), "encrypted": encrypted, "text_length": 0}
    return records, "", metrics, warnings, {"container_type": path.suffix.lower()}


def extract_email(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    if path.suffix.lower() != ".eml":
        return [], "", {"text_length": 0}, [warning("FORMATO_NO_SOPORTADO_COMPLETAMENTE", "MSG requiere un extractor estático adicional.")], {}
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    bodies = []
    attachments = []
    for part in message.walk():
        disposition = part.get_content_disposition()
        if disposition == "attachment":
            payload = part.get_payload(decode=True) or b""
            attachments.append({
                "record_type": "attachment", "filename": part.get_filename(), "content_type": part.get_content_type(),
                "size_bytes": len(payload), "sha256": hashlib.sha256(payload).hexdigest(),
            })
        elif part.get_content_type() == "text/plain":
            try:
                bodies.append(part.get_content())
            except Exception:
                pass
    body = "\n".join(bodies)[:MAX_TEXT_CHARS]
    record = {
        "record_type": "email_message", "from": str(message.get("From", "")),
        "to": str(message.get("To", "")), "cc": str(message.get("Cc", "")),
        "date": str(message.get("Date", "")), "subject": str(message.get("Subject", "")), "body": body,
    }
    warnings = [warning("POSIBLE_INFORMACION_PERSONAL", "El correo puede contener datos personales; revisar antes de difundir.", "ALTO")]
    metrics = {"attachment_count": len(attachments), "text_length": len(body)}
    return [record, *attachments], body, metrics, warnings, {"message_id": str(message.get("Message-ID", ""))}


def extract_sqlite(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    records: list[dict[str, Any]] = []
    warnings = []
    uri = f"file:{quote(str(path.resolve()))}?mode=ro"
    connection = sqlite3.connect(uri, uri=True)
    try:
        tables = connection.execute(
            "SELECT name, sql FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
        ).fetchall()
        for name, sql in tables:
            try:
                count = connection.execute(f"SELECT COUNT(*) FROM {quote_sql_identifier(name)}").fetchone()[0]
            except sqlite3.Error:
                count = None
            records.append({"record_type": "database_table", "name": name, "row_count": count, "definition": sql})
            for cid, field, data_type, not_null, default, primary_key in connection.execute(f"PRAGMA table_info({quote_sql_identifier(name)})"):
                records.append({
                    "record_type": "database_field", "table": name, "ordinal": cid, "name": field,
                    "data_type": data_type, "not_null": bool(not_null), "default": default, "primary_key": bool(primary_key),
                })
            for relation in connection.execute(f"PRAGMA foreign_key_list({quote_sql_identifier(name)})"):
                records.append({
                    "record_type": "database_relation", "table": name, "target_table": relation[2],
                    "source_field": relation[3], "target_field": relation[4],
                })
    finally:
        connection.close()
    metrics = {"database_table_count": sum(r["record_type"] == "database_table" for r in records), "text_length": 0}
    return records, "", metrics, warnings, {"engine": "SQLite", "access_mode": "READ_ONLY"}


def extract_code(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    raw = path.read_bytes()
    text = raw[:MAX_TEXT_CHARS].decode("utf-8", errors="replace")
    records: list[dict[str, Any]] = [{"record_type": "code_file", "language": path.suffix.lower().lstrip("."), "text": text}]
    dependencies: set[str] = set()
    symbols: list[tuple[str, str, int]] = []
    if path.suffix.lower() == ".py":
        try:
            tree = ast.parse(text)
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    symbols.append((node.name, "FUNCTION", node.lineno))
                elif isinstance(node, ast.ClassDef):
                    symbols.append((node.name, "CLASS", node.lineno))
                elif isinstance(node, ast.Import):
                    dependencies.update(alias.name for alias in node.names)
                elif isinstance(node, ast.ImportFrom) and node.module:
                    dependencies.add(node.module)
        except SyntaxError:
            pass
    else:
        for match in re.finditer(r"(?m)^\s*(?:class|function|def|fn)\s+([A-Za-z_$][\w$]*)", text):
            symbols.append((match.group(1), "SYMBOL", text.count("\n", 0, match.start()) + 1))
        for match in re.finditer(r"(?m)^\s*(?:import|require\s*\()\s*['\"]?([^'\";\s)]+)", text):
            dependencies.add(match.group(1))
    records.extend({"record_type": "code_symbol", "name": name, "symbol_type": kind, "line": line} for name, kind, line in symbols)
    records.extend({"record_type": "dependency", "name": value, "evidence": "STATIC_PARSE"} for value in sorted(dependencies))
    secret_pattern = re.compile(r"(?i)(api[_-]?key|secret|token|password)\s*[:=]")
    warnings = []
    if secret_pattern.search(text):
        warnings.append(warning("SECRETO_APARENTE", "Se detectó un posible secreto por análisis estático; su valor no se reproduce en el resumen.", "ALTO"))
    if len(raw) > MAX_TEXT_CHARS:
        warnings.append(warning("EXTRACCION_PARCIAL", f"Lectura textual limitada a {MAX_TEXT_CHARS} bytes."))
    metrics = {"code_symbol_count": len(symbols), "dependency_count": len(dependencies), "text_length": len(text)}
    return records, text, metrics, warnings, {"language": path.suffix.lower().lstrip("."), "parsed_without_execution": True}


def extract_textual(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    raw = path.read_bytes()
    text = raw[:MAX_TEXT_CHARS].decode("utf-8", errors="replace")
    warnings = []
    if len(raw) > MAX_TEXT_CHARS:
        warnings.append(warning("EXTRACCION_PARCIAL", f"Lectura textual limitada a {MAX_TEXT_CHARS} bytes."))
    records = [{"record_type": "content", "text": text}]
    return records, text, {"text_length": len(text)}, warnings, {}


def extract_csv(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    raw = path.read_bytes()
    text = raw[:MAX_TEXT_CHARS].decode("utf-8-sig", errors="replace")
    records: list[dict[str, Any]] = [{"record_type": "sheet", "sheet_id": "SHT-001", "name": path.stem, "order": 1, "visibility": "VISIBLE"}]
    row_count = cell_count = 0
    for row_number, row in enumerate(csv.reader(text.splitlines()), 1):
        row_count += 1
        for column_number, value in enumerate(row, 1):
            if value == "":
                continue
            cell_count += 1
            records.append({
                "record_type": "cell", "sheet_id": "SHT-001", "row": row_number,
                "column": column_number, "value": value, "data_type": "TEXT",
            })
    warnings = []
    if len(raw) > MAX_TEXT_CHARS:
        warnings.append(warning("EXTRACCION_PARCIAL", f"Lectura CSV limitada a {MAX_TEXT_CHARS} bytes."))
    return records, text, {"sheet_count": 1, "sheet_names": [path.stem], "row_count": row_count, "cell_count": cell_count, "text_length": len(text)}, warnings, {}


def extract_technical(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    ext = path.suffix.lower()
    records: list[dict[str, Any]] = []
    text = ""
    metrics: dict[str, Any] = {"text_length": 0}
    warnings = []
    if ext == ".geojson":
        data = json.loads(path.read_text(encoding="utf-8"))
        features = data.get("features", []) if isinstance(data, dict) else []
        for index, feature in enumerate(features, 1):
            geometry = feature.get("geometry") or {}
            records.append({
                "record_type": "gis_feature", "feature_number": index, "feature_id": feature.get("id"),
                "geometry_type": geometry.get("type"), "properties": json_safe(feature.get("properties", {})),
            })
        metrics.update({"entity_count": len(features), "gis_feature_count": len(features)})
    elif ext == ".ifc":
        text = path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS]
        counts = Counter(re.findall(r"=\s*(IFC[A-Z0-9_]+)\s*\(", text, re.I))
        records.extend({"record_type": "entity", "entity_type": kind.upper(), "count": count} for kind, count in sorted(counts.items()))
        metrics.update({"entity_count": sum(counts.values()), "text_length": len(text)})
        warnings.append(warning("GEOMETRIA_NO_REPRESENTADA", "Se inventariaron entidades IFC, pero no se reprodujo toda la geometría."))
    elif ext == ".dxf":
        text = path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS]
        entity_types = Counter(re.findall(r"(?m)^\s*0\s*\n\s*([A-Z][A-Z0-9_]*)\s*$", text))
        records.extend({"record_type": "entity", "entity_type": kind, "count": count} for kind, count in sorted(entity_types.items()))
        metrics.update({"entity_count": sum(entity_types.values()), "text_length": len(text)})
        warnings.append(warning("GEOMETRIA_NO_REPRESENTADA", "Se inventariaron tipos DXF, pero no se reprodujo toda la geometría."))
    else:
        warnings.append(warning("FORMATO_NO_SOPORTADO_COMPLETAMENTE", "El formato técnico se verificó por metadatos y hash; no se interpretó su modelo interno."))
        warnings.append(warning("GEOMETRIA_NO_REPRESENTADA", "No se reprodujo la geometría del modelo.", "BAJO"))
    return records, text, metrics, warnings, {"software_family": technical_family(ext), "parsed_without_execution": True}


def technical_family(ext: str) -> str | None:
    if ext in {".edb", ".e2k"}:
        return "ETABS"
    if ext in {".sdb", ".f2k"}:
        return "SAP2000"
    if ext in {".rvt", ".rfa"}:
        return "REVIT"
    if ext in {".dwg", ".dxf", ".dwt", ".dwf", ".dwfx"}:
        return "AUTOCAD"
    if ext in {".qgz", ".qgs", ".shp", ".geojson", ".kml", ".kmz"}:
        return "GIS"
    return None


def extract_static_binary(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    metadata: dict[str, Any] = {"executed": False, "loaded": False}
    try:
        completed = subprocess.run(["file", "--brief", "--", str(path)], capture_output=True, text=True, timeout=15, check=True)
        metadata["static_signature_description"] = completed.stdout.strip()
    except Exception as exc:
        metadata["static_signature_error"] = type(exc).__name__
    warnings = [warning("CONTENIDO_ACTIVO_NO_EJECUTADO", "El archivo activo fue inspeccionado solo de forma estática; no se ejecutó, instaló ni cargó.", "ALTO")]
    return [], "", {"text_length": 0}, warnings, metadata


def extract_generic(path: Path) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    return [], "", {"text_length": 0}, [warning("FORMATO_NO_SOPORTADO_COMPLETAMENTE", "Se conservaron identidad, integridad y metadatos; no existe extractor determinista completo.")], {}


def deterministic_extract(path: Path, role: str) -> tuple[list[dict[str, Any]], str, dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    ext = path.suffix.lower()
    try:
        if role == "EJECUTABLE_ACTIVO":
            return extract_static_binary(path)
        if role == "CONTENEDOR":
            return extract_archive(path)
        if role == "BASE_DATOS" and ext in {".sqlite", ".sqlite3", ".db"}:
            return extract_sqlite(path)
        if role == "CODIGO_FUENTE":
            return extract_code(path)
        if ext == ".pdf":
            return extract_pdf(path)
        if ext in {".docx", ".docm"}:
            return extract_docx(path)
        if ext in {".xlsx", ".xlsm"}:
            return extract_xlsx(path)
        if ext == ".csv":
            return extract_csv(path)
        if ext in {".pptx", ".pptm"}:
            return extract_presentation(path)
        if ext in IMAGES:
            return extract_image(path)
        if ext in AUDIO:
            return extract_media(path, "audio")
        if ext in VIDEO:
            return extract_media(path, "video")
        if ext in EMAILS:
            return extract_email(path)
        if ext in TECHNICAL:
            return extract_technical(path)
        if ext in {".txt", ".md", ".rtf", ".log", ".xml", ".json", ".yaml", ".yml", ".ini", ".cfg", ".sql"}:
            return extract_textual(path)
        return extract_generic(path)
    except Exception as exc:
        return [], "", {"text_length": 0}, [warning("EXTRACCION_PARCIAL", f"El extractor falló de forma controlada: {type(exc).__name__}: {exc}", "ALTO")], {}


def normalized_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def normalized_title(stem: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", stem.casefold()).strip()


def keyword_values(path: Path, text: str, metadata: dict[str, Any]) -> list[str]:
    keyword_text = redact_sensitive_text(text[:200000])
    tokens = re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ][\wÁÉÍÓÚÜÑáéíóúüñ-]{2,}", f"{path.stem} {keyword_text}")
    counts = Counter(token.casefold() for token in tokens if token.casefold() not in STOPWORDS and not token.isdigit())
    values = [token for token, _ in counts.most_common(24)]
    family = metadata.get("software_family")
    if family and family.casefold() not in values:
        values.append(family.casefold())
    return values


def redact_sensitive_text(text: str) -> str:
    return re.sub(
        r"(?i)(api[_-]?key|secret|token|password)\s*[:=]\s*(['\"]?)[^\s,'\";]+\2",
        r"\1 REDACTED",
        text,
    )


def summary_text(path: Path, role: str, detected: dict[str, Any], metrics: dict[str, Any], warnings: list[dict[str, Any]]) -> str:
    parts = [f"{path.name} es un recurso {detected['declared_format']} con rol operativo {role}."]
    observed = []
    for key in ("page_count", "sheet_count", "slide_count", "table_count", "formula_count", "image_count", "archive_member_count", "database_table_count", "entity_count", "attachment_count"):
        if metrics.get(key) is not None:
            observed.append(f"{key}={metrics[key]}")
    if observed:
        parts.append("Componentes observados: " + ", ".join(observed) + ".")
    parts.append("Puede utilizarse para consulta y comparación posterior; la clasificación definitiva corresponde a otra etapa.")
    if warnings:
        parts.append("La extracción presenta limitaciones registradas explícitamente en el JSONL.")
    return " ".join(parts)


def deterministic_descriptions(
    path: Path,
    role: str,
    detected: dict[str, Any],
    metrics: dict[str, Any],
    warnings: list[dict[str, Any]],
) -> dict[str, Any]:
    brief = summary_text(path, role, detected, metrics, warnings)
    metric_names = [
        key for key in (
            "page_count", "sheet_count", "slide_count", "table_count", "formula_count",
            "image_count", "archive_member_count", "database_table_count", "entity_count",
            "attachment_count",
        ) if metrics.get(key) is not None
    ]
    observed = ", ".join(metric_names) if metric_names else "metadatos, formato e integridad"
    long_description = (
        f"{brief} La representación de máquina conserva {observed} cuando fueron verificables. "
        "Los hechos objetivos proceden de extracción determinista; las limitaciones y el alcance "
        "quedan registrados sin sustituir al archivo original."
    )
    return {
        "short_description": f"{path.stem}: recurso {detected['declared_format']} ({role}).",
        "brief_description": brief,
        "long_description": long_description,
        "synonyms": [],
    }


def structural_signature(records: Iterable[dict[str, Any]], metrics: dict[str, Any]) -> str:
    structure = {
        "record_counts": Counter(record["record_type"] for record in records),
        "metrics": metrics,
    }
    payload = json.dumps(json_safe(structure), ensure_ascii=False, sort_keys=True).encode()
    return hashlib.sha256(payload).hexdigest()


def build_records(
    path: Path,
    root: Path,
    role: str,
    role_basis: str,
    relations: list[dict[str, Any]] | None = None,
    enrichment: dict[str, Any] | None = None,
    processing_mode: str = PROCESSING_MODE,
    processing_metadata: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    source_hash = sha256_file(path)
    stat = path.stat()
    detected = detect_format(path)
    content_records, text, metrics, extractor_warnings, technical_metadata = deterministic_extract(path, role)
    source_metadata = base_metadata(path, root)
    source_metadata.update({
        "format_declared": detected["declared_format"],
        "format_detected": detected["detected_format"],
        "title_detected": metadata_value(technical_metadata, "title"),
        "author_detected": metadata_value(technical_metadata, "author", "creator"),
        "organization_detected": metadata_value(technical_metadata, "organization", "company"),
        "creator_application": metadata_value(technical_metadata, "creator_application", "application", "producer", "creator_tool"),
        "modifier_application": metadata_value(technical_metadata, "modifier_application", "last_modified_by"),
    })
    generated_at = utc_now()
    ai_applied = bool(enrichment)
    manifest = {
        "record_type": "manifest", "schema_name": SCHEMA_NAME, "schema_version": SCHEMA_VERSION,
        "source_filename": path.name, "source_relative_path": source_metadata["source_relative_path"],
        "generated_at": generated_at,
        "generated_by": f"{SKILL_NAME} {SKILL_VERSION}" + (" + IA_UNICA" if ai_applied else ""),
        "processing_mode": processing_mode,
    }
    manifest.update(json_safe(processing_metadata or {}))
    integrity = {"record_type": "integrity", "algorithm": "SHA-256", "source_sha256": source_hash, "source_size_bytes": stat.st_size}
    role_record = {"record_type": "file_role", "role": role, "basis": role_basis, "classification_is_final": False}
    format_record = {"record_type": "format_information", **detected, **json_safe(technical_metadata)}
    normalized = normalized_text(text)
    comparison = {
        "record_type": "comparison_features", "source_sha256": source_hash,
        "source_size_bytes": stat.st_size, "source_created_at": source_metadata["source_created_at"],
        "source_modified_at": source_metadata["source_modified_at"], "detected_format": detected["detected_format"],
        "page_count": metrics.get("page_count"), "sheet_count": metrics.get("sheet_count"),
        "sheet_names": metrics.get("sheet_names"), "table_count": metrics.get("table_count"),
        "formula_count": metrics.get("formula_count"), "image_count": metrics.get("image_count"),
        "layer_count": metrics.get("layer_count"), "entity_count": metrics.get("entity_count"),
        "text_length": len(text), "normalized_title": normalized_title(path.stem),
        "normalized_text_hash": hashlib.sha256(normalized.encode()).hexdigest() if normalized else None,
        "structure_signature": structural_signature(content_records, metrics),
        "metadata_signature": hashlib.sha256(json.dumps(json_safe({**detected, **technical_metadata}), sort_keys=True).encode()).hexdigest(),
    }
    descriptions = deterministic_descriptions(path, role, detected, metrics, extractor_warnings)
    keywords = keyword_values(path, text, technical_metadata)
    semantic_records: list[dict[str, Any]] = []
    if enrichment:
        legacy_summary = enrichment.get("summary")
        if isinstance(legacy_summary, str) and legacy_summary.strip():
            descriptions["brief_description"] = redact_sensitive_text(legacy_summary.strip())
        for key in ("short_description", "brief_description", "long_description"):
            if isinstance(enrichment.get(key), str) and enrichment[key].strip():
                descriptions[key] = redact_sensitive_text(enrichment[key].strip())
        if isinstance(enrichment.get("synonyms"), list):
            descriptions["synonyms"] = sorted({
                redact_sensitive_text(str(item).strip())
                for item in enrichment["synonyms"] if str(item).strip()
            })
        if isinstance(enrichment.get("keywords"), list):
            keywords = sorted(set(keywords + [
                redact_sensitive_text(str(item).strip())
                for item in enrichment["keywords"] if str(item).strip()
            ]))
        for item in enrichment.get("entities", []):
            if isinstance(item, dict) and item.get("name"):
                if not item.get("evidence"):
                    raise ValueError(f"Entidad IA sin evidencia para {path.name}: {item.get('name')}")
                semantic_records.append({"record_type": "entity", **json_safe(item), "evidence": item.get("evidence")})
        for item in enrichment.get("relations", []):
            if isinstance(item, dict) and item.get("relation_type"):
                if not item.get("evidence"):
                    raise ValueError(f"Relación IA sin evidencia para {path.name}: {item.get('relation_type')}")
                semantic_records.append({"record_type": "relation", **json_safe(item), "status": item.get("status", "OBSERVADA")})
    all_warnings = list(extractor_warnings)
    if not ai_applied:
        all_warnings.append(warning("IA_NO_APLICADA", "La salida contiene extracción y síntesis deterministas; la intervención semántica de IA no estuvo disponible.", "BAJO"))
    security = {
        "record_type": "security", "source_executed": False, "macros_executed": False,
        "scripts_executed": False, "active_content_detected": role == "EJECUTABLE_ACTIVO" or any(item["code"] in {"MACRO_NO_EJECUTADA", "CONTENIDO_ACTIVO_NO_EJECUTADO"} for item in all_warnings),
        "possible_secret_detected": any(item["code"] == "SECRETO_APARENTE" for item in all_warnings),
    }
    rights_value = metadata_value(technical_metadata, "copyright", "rights", "license")
    confidence = 0.98
    if any(item["code"] in {"FORMATO_NO_SOPORTADO_COMPLETAMENTE", "EXTRACCION_PARCIAL"} for item in all_warnings):
        confidence = 0.65
    elif any(item["code"] == "NO_EXTRAIDO" for item in all_warnings):
        confidence = 0.78
    quality = {
        "record_type": "quality", "extraction_confidence": confidence,
        "extraction_status": "PARTIAL" if confidence < 0.9 else "COMPLETE",
        "deterministic_extraction": True, "ai_enrichment_applied": ai_applied,
        "ai_passes": 1 if ai_applied else 0, "human_review_recommended": bool(all_warnings),
    }
    relation_records = [{"record_type": "relation", **json_safe(item)} for item in (relations or [])]
    records = [manifest, source_metadata, integrity, role_record, format_record]
    records.extend(content_records)
    if metrics:
        records.append({"record_type": "technical_data", **json_safe(metrics)})
    records.extend(semantic_records)
    records.extend(relation_records)
    records.extend([
        {
            "record_type": "summary",
            "text": descriptions["brief_description"],
            **descriptions,
            "generation_method": "AI_UNICA" if ai_applied else "DETERMINISTIC",
        },
        {"record_type": "keywords", "values": keywords, "classification_is_final": False},
    ])
    if rights_value:
        records.append({"record_type": "rights_observed", "value": str(rights_value), "status": "OBSERVED_NOT_VERIFIED"})
    records.extend([security, comparison, quality])
    records.extend(all_warnings)
    records.extend([
        {"record_type": "warnings", "codes": [item["code"] for item in all_warnings], "count": len(all_warnings)},
        {"record_type": "end", "status": "COMPLETED_WITH_WARNINGS" if all_warnings else "COMPLETED", "source_sha256": source_hash},
    ])
    return sorted(records, key=lambda record: ORDER.get(record["record_type"], 7))


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    path.write_text("".join(json.dumps(record, ensure_ascii=False, separators=(",", ":"), allow_nan=False) + "\n" for record in records), encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    records = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            raise ValueError(f"Línea vacía en JSONL: {path}:{line_number}")
        value = json.loads(line)
        if not isinstance(value, dict) or not isinstance(value.get("record_type"), str):
            raise ValueError(f"Registro JSONL inválido: {path}:{line_number}")
        records.append(value)
    return records


def validate_jsonl(
    path: Path,
    source: Path,
    expected_mode: str = PROCESSING_MODE,
) -> list[dict[str, Any]]:
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError(f"JSONL ausente o vacío: {path}")
    records = read_jsonl(path)
    for line_number, record in enumerate(records, 1):
        record_type = record["record_type"]
        if record_type not in ORDER:
            raise ValueError(f"record_type no admitido en línea {line_number}: {record_type}")
        missing_fields = REQUIRED_BY_TYPE.get(record_type, set()).difference(record)
        if missing_fields:
            raise ValueError(f"Campos ausentes en {record_type}: {sorted(missing_fields)}")
    ranks = [ORDER.get(record["record_type"], 7) for record in records]
    if ranks != sorted(ranks):
        raise ValueError(f"Jerarquía JSONL fuera de orden: {path}")
    record_types = Counter(record["record_type"] for record in records)
    missing = UNIVERSAL.difference(record_types)
    if missing:
        raise ValueError(f"Registros universales ausentes: {sorted(missing)}")
    for unique in UNIVERSAL:
        if record_types[unique] != 1:
            raise ValueError(f"Registro universal repetido: {unique}")
    manifest = next(record for record in records if record["record_type"] == "manifest")
    source_metadata = next(record for record in records if record["record_type"] == "source_metadata")
    integrity = next(record for record in records if record["record_type"] == "integrity")
    role = next(record for record in records if record["record_type"] == "file_role")
    summary = next(record for record in records if record["record_type"] == "summary")
    keywords = next(record for record in records if record["record_type"] == "keywords")
    quality = next(record for record in records if record["record_type"] == "quality")
    warnings_record = next(record for record in records if record["record_type"] == "warnings")
    end = records[-1]
    source_hash = sha256_file(source)
    if (
        manifest.get("schema_name") != SCHEMA_NAME
        or manifest.get("schema_version") != SCHEMA_VERSION
        or manifest.get("processing_mode") != expected_mode
        or manifest.get("source_filename") != source.name
    ):
        raise ValueError("Manifest no corresponde al modo esperado o a la fuente.")
    if (
        integrity.get("algorithm") != "SHA-256"
        or not SHA256_RE.fullmatch(str(integrity.get("source_sha256", "")))
        or integrity.get("source_sha256") != source_hash
        or integrity.get("source_size_bytes") != source.stat().st_size
    ):
        raise ValueError("Integridad del JSONL no corresponde a la fuente.")
    if (
        source_metadata.get("source_filename") != source.name
        or source_metadata.get("source_size_bytes") != source.stat().st_size
        or source_metadata.get("source_relative_path") != manifest.get("source_relative_path")
    ):
        raise ValueError("Metadatos universales no corresponden a la fuente o al manifest.")
    if role.get("role") not in ALLOWED_ROLES or role.get("classification_is_final") is not False:
        raise ValueError("Rol operativo inválido o presentado como clasificación final.")
    description_fields = ("short_description", "brief_description", "long_description")
    if any(not isinstance(summary.get(key), str) or not summary[key].strip() for key in description_fields):
        raise ValueError("Descripción corta, breve o larga ausente o vacía.")
    if summary.get("text") != summary.get("brief_description"):
        raise ValueError("El alias summary.text debe coincidir con brief_description.")
    if not isinstance(summary.get("synonyms"), list) or any(not isinstance(item, str) or not item.strip() for item in summary["synonyms"]):
        raise ValueError("Lista de sinónimos inválida.")
    if not isinstance(keywords.get("values"), list) or keywords.get("classification_is_final") is not False:
        raise ValueError("Registro de palabras clave inválido.")
    confidence = quality.get("extraction_confidence")
    if not isinstance(confidence, (int, float)) or isinstance(confidence, bool) or not 0 <= confidence <= 1:
        raise ValueError("Confianza de extracción inválida.")
    if quality.get("ai_passes") not in {0, 1}:
        raise ValueError("La IA solo puede intervenir cero o una vez.")
    warning_codes = [record["code"] for record in records if record["record_type"] == "warning"]
    if warnings_record.get("codes") != warning_codes or warnings_record.get("count") != len(warning_codes):
        raise ValueError("Agregado de advertencias inconsistente.")
    if end.get("record_type") != "end" or end.get("source_sha256") != source_hash:
        raise ValueError("Registro END no corresponde a la fuente.")
    return records


def find_record(records: list[dict[str, Any]], record_type: str) -> dict[str, Any]:
    return next(record for record in records if record["record_type"] == record_type)


def render_markdown(records: list[dict[str, Any]], jsonl_hash: str) -> str:
    manifest = find_record(records, "manifest")
    metadata = find_record(records, "source_metadata")
    integrity = find_record(records, "integrity")
    role = find_record(records, "file_role")
    fmt = find_record(records, "format_information")
    summary = find_record(records, "summary")
    keywords = find_record(records, "keywords")
    quality = find_record(records, "quality")
    warnings = [record for record in records if record["record_type"] == "warning"]
    relations = [record for record in records if record["record_type"] == "relation"]
    technical = next((record for record in records if record["record_type"] == "technical_data"), {})
    title = metadata.get("title_detected") or Path(manifest["source_filename"]).stem
    lines = [
        "---", f"schema_name: {SCHEMA_NAME}", f"schema_version: {SCHEMA_VERSION}",
        f"source_filename: {json.dumps(manifest['source_filename'], ensure_ascii=False)}",
        f"source_relative_path: {json.dumps(manifest['source_relative_path'], ensure_ascii=False)}",
        f"source_size_bytes: {metadata['source_size_bytes']}",
        f"source_sha256: {integrity['source_sha256']}", f"jsonl_sha256: {jsonl_hash}", f"file_role: {role['role']}",
        f"format_detected: {json.dumps(fmt.get('detected_format'), ensure_ascii=False)}",
        f"title_detected: {json.dumps(title, ensure_ascii=False)}", f"language: {json.dumps(metadata.get('language'))}",
        f"extraction_status: {quality['extraction_status']}", f"extraction_confidence: {quality['extraction_confidence']}",
        f"generated_at: {manifest['generated_at']}", "---", "", f"# {title}", "",
        "## Descripción corta", "", summary["short_description"], "",
        "## Descripción breve", "", summary["brief_description"], "",
        "## Descripción larga", "", summary["long_description"], "",
        "## Contenido principal", "",
        f"- Recurso: `{manifest['source_filename']}`", f"- Rol operativo: `{role['role']}`",
        f"- Formato detectado: `{fmt.get('detected_format')}`", "", "## Estructura", "",
    ]
    if technical:
        for key, value in technical.items():
            if key != "record_type" and value is not None:
                lines.append(f"- {key}: `{json.dumps(value, ensure_ascii=False)}`")
    else:
        lines.append("- Sin estructura adicional extraíble de forma determinista.")
    lines.extend(["", "## Datos técnicos relevantes", "", f"- Tamaño: `{metadata['source_size_bytes']}` bytes", f"- SHA-256: `{integrity['source_sha256']}`", "", "## Sinónimos", ""])
    lines.extend(f"- {value}" for value in summary.get("synonyms", []))
    if not summary.get("synonyms"):
        lines.append("- Sin sinónimos sustentados.")
    lines.extend(["", "## Palabras clave", ""])
    lines.extend(f"- {value}" for value in keywords.get("values", []))
    if not keywords.get("values"):
        lines.append("- Sin palabras clave verificables.")
    lines.extend(["", "## Relaciones observadas", ""])
    if relations:
        lines.extend(f"- {item.get('relation_type')}: {item.get('target_filename')} ({item.get('status', 'OBSERVADA')})" for item in relations)
    else:
        lines.append("- No se confirmaron relaciones automáticas.")
    lines.extend(["", "## Limitaciones de extracción", ""])
    if warnings:
        lines.extend(f"- `{item['code']}`: {item['description']}" for item in warnings)
    else:
        lines.append("- Sin limitaciones registradas.")
    lines.extend([
        "", "## Calidad", "", f"- Cobertura: `{quality['extraction_status']}`",
        f"- Confianza: `{quality['extraction_confidence']}`", f"- Advertencias: `{len(warnings)}`",
        f"- Revisión humana: `{'RECOMENDADA' if quality['human_review_recommended'] else 'OPCIONAL'}`",
        "", "## Próxima etapa", "",
        "El recurso está preparado para un proceso posterior de clasificación e incorporación a ECOSISTEMA. Este derivado no decide su destino definitivo.", "",
    ])
    return "\n".join(lines)


def sha256_text(source: Path) -> str:
    source_hash = sha256_file(source)
    return "\n".join([
        f"# schema_name: {SCHEMA_NAME}",
        f"# schema_version: {SCHEMA_VERSION}",
        "# artifact_type: integrity",
        f"# source_filename: {json.dumps(source.name, ensure_ascii=False)}",
        f"# source_size_bytes: {source.stat().st_size}",
        f"# source_sha256: {source_hash}",
        f"{source_hash}  {source.name}",
        "",
    ])


def wants_markdown(role: str, records: list[dict[str, Any]]) -> bool:
    if role in {"FUENTE_PRIMARIA", "MANIFESTACION_ASOCIADA", "BASE_DATOS", "CODIGO_FUENTE"}:
        return True
    if role in {"ARCHIVO_AUXILIAR", "DERIVADO_GENERADO"}:
        comparison = find_record(records, "comparison_features")
        return bool(comparison.get("text_length"))
    return False


def output_paths(source: Path) -> dict[str, Path]:
    return {
        "jsonl": source.with_name(source.name + ".jsonl"),
        "markdown": source.with_name(source.name + ".md"),
        "sha256": source.with_name(source.name + ".sha256"),
    }


def validate_bundle(stage: Path, source: Path, include_jsonl: bool, include_markdown: bool) -> None:
    paths = output_paths(source)
    staged_jsonl = stage / paths["jsonl"].name
    staged_md = stage / paths["markdown"].name
    staged_sha = stage / paths["sha256"].name
    if not staged_sha.is_file() or staged_sha.read_text(encoding="utf-8") != sha256_text(source):
        raise ValueError("Archivo SHA-256 inválido o no corresponde al original.")
    if include_jsonl:
        records = validate_jsonl(staged_jsonl, source)
        if include_markdown:
            expected = render_markdown(records, sha256_file(staged_jsonl))
            if staged_md.read_text(encoding="utf-8") != expected:
                raise ValueError("El Markdown no fue generado exclusivamente desde el JSONL validado.")
        elif staged_md.exists():
            raise ValueError("Se generó Markdown no autorizado para este rol.")
    else:
        if staged_jsonl.exists() or staged_md.exists():
            raise ValueError("El rol solo autoriza SHA-256.")


def publish_bundle(stage: Path, source: Path, desired: set[str]) -> None:
    paths = output_paths(source)
    managed = {key: value for key, value in paths.items()}
    for destination in managed.values():
        if destination.is_symlink() or (destination.exists() and not destination.is_file()):
            raise ValueError(f"Colisión insegura en derivado lateral: {destination}")
    backup = Path(tempfile.mkdtemp(prefix=f".{source.name}.lhm-backup-", dir=source.parent))
    moved_old: list[tuple[Path, Path]] = []
    published: list[Path] = []
    keep_backup = False
    try:
        for key, destination in managed.items():
            if destination.exists():
                old = backup / destination.name
                os.replace(destination, old)
                moved_old.append((old, destination))
        for key in ("jsonl", "markdown", "sha256"):
            if key not in desired:
                continue
            staged = stage / paths[key].name
            os.replace(staged, paths[key])
            published.append(paths[key])
    except Exception as publish_error:
        for destination in published:
            if destination.exists():
                destination.unlink()
        restore_errors = []
        for old, destination in reversed(moved_old):
            if old.exists():
                try:
                    os.replace(old, destination)
                except Exception as restore_error:
                    restore_errors.append(f"{destination}: {restore_error}")
        if restore_errors:
            keep_backup = True
            raise RuntimeError(
                "Falló la publicación lateral y la restauración no fue completa; "
                f"los derivados recuperables permanecen en {backup}: {'; '.join(restore_errors)}"
            ) from publish_error
        raise publish_error
    finally:
        if not keep_backup:
            shutil.rmtree(backup, ignore_errors=True)


def load_enrichment(path: Path | None) -> dict[str, Any]:
    if path is None:
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("El enriquecimiento IA debe ser un objeto por ruta relativa.")
    return value


def prepare_descriptors(target: Path) -> tuple[Path, list[dict[str, Any]]]:
    root = target if target.is_dir() else target.parent
    descriptors = []
    for path in iter_sources(target):
        role, basis = preliminary_role(path)
        digest = sha256_file(path)
        descriptors.append({"path": path, "hash": digest, "size": path.stat().st_size, "role": role, "basis": basis, "stem": path.stem, "extension": path.suffix.lower()})
    by_hash: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in descriptors:
        by_hash[item["hash"]].append(item)
    for items in by_hash.values():
        if len(items) > 1:
            for duplicate in sorted(items, key=lambda item: item["path"].relative_to(root).as_posix())[1:]:
                duplicate["role"] = "RESPALDO_COPIA"
                duplicate["basis"] = "DUPLICADO_EXACTO_SHA256"
    by_stem: dict[tuple[Path, str], list[dict[str, Any]]] = defaultdict(list)
    for item in descriptors:
        by_stem[(item["path"].parent, item["stem"])].append(item)
    for items in by_stem.values():
        technical = next((item for item in items if item["extension"] in TECHNICAL_PRIMARY and item["role"] == "FUENTE_PRIMARIA"), None)
        if technical:
            for item in items:
                if item is not technical and item["extension"] in GENERATED_CANDIDATES and item["role"] == "FUENTE_PRIMARIA":
                    item["role"] = "DERIVADO_GENERADO"
                    item["basis"] = f"MISMO_NOMBRE_BASE_QUE_{technical['path'].name}"
                    item["relation"] = {"relation_type": "GENERADO_POR", "target_filename": technical["path"].name, "status": "PROBABLE"}
                elif item["role"] == "ARCHIVO_AUXILIAR":
                    item["relation"] = {"relation_type": "DEPENDE_DE", "target_filename": technical["path"].name, "status": "PROBABLE"}
                elif item["role"] == "RESPALDO_COPIA":
                    item["relation"] = {"relation_type": "RESPALDO_DE", "target_filename": technical["path"].name, "status": "PROBABLE"}
        else:
            candidates = sorted(
                (item for item in items if item["role"] == "FUENTE_PRIMARIA"),
                key=lambda item: item["path"].name,
            )
            if len(candidates) > 1:
                text_hashes: dict[str, list[dict[str, Any]]] = defaultdict(list)
                for item in candidates:
                    _, text_value, _, _, _ = deterministic_extract(item["path"], item["role"])
                    normalized = normalized_text(text_value)
                    if normalized:
                        text_hashes[hashlib.sha256(normalized.encode()).hexdigest()].append(item)
                associated_ids: set[int] = set()
                for matching in text_hashes.values():
                    if len(matching) < 2:
                        continue
                    canonical = matching[0]
                    for item in matching[1:]:
                        item["role"] = "MANIFESTACION_ASOCIADA"
                        item["basis"] = "NOMBRE_BASE_EXACTO_Y_HASH_TEXTO_NORMALIZADO_COINCIDENTE"
                        item["relation"] = {
                            "relation_type": "MANIFESTACION_DE", "target_filename": canonical["path"].name,
                            "status": "CONFIRMADA", "evidence": "NORMALIZED_TEXT_HASH_MATCH",
                        }
                        associated_ids.add(id(item))
                canonical = candidates[0]
                for item in candidates[1:]:
                    if id(item) not in associated_ids:
                        item["relation"] = {
                            "relation_type": "RELACIONADO_CON", "target_filename": canonical["path"].name,
                            "status": "POSIBLE", "evidence": "NOMBRE_BASE_EXACTO_SIN_EVIDENCIA_DE_CONTENIDO_SUFFICIENTE",
                        }
    return root, descriptors


def process_one(item: dict[str, Any], root: Path, enrichment: dict[str, Any], forced_failure: bool = False) -> dict[str, Any]:
    source: Path = item["path"]
    original_hash = item["hash"]
    role = item["role"]
    if role == "TEMPORAL_CACHE":
        return {"source": str(source), "role": role, "status": "DETECTADO_SIN_DERIVADOS"}
    include_jsonl = role != "RESPALDO_COPIA"
    stage = Path(tempfile.mkdtemp(prefix=f".{source.name}.lhm-stage-", dir=source.parent))
    try:
        paths = output_paths(source)
        desired = {"sha256"}
        (stage / paths["sha256"].name).write_text(sha256_text(source), encoding="utf-8")
        include_md = False
        if include_jsonl:
            relative = source.relative_to(root).as_posix() if root.is_dir() else source.name
            records = build_records(source, root, role, item["basis"], [item["relation"]] if item.get("relation") else [], enrichment.get(relative))
            staged_jsonl = stage / paths["jsonl"].name
            write_jsonl(staged_jsonl, records)
            validate_jsonl(staged_jsonl, source)
            desired.add("jsonl")
            include_md = wants_markdown(role, records)
            if include_md:
                (stage / paths["markdown"].name).write_text(render_markdown(records, sha256_file(staged_jsonl)), encoding="utf-8")
                desired.add("markdown")
        if forced_failure:
            raise ValueError("Fallo de validación solicitado para prueba.")
        validate_bundle(stage, source, include_jsonl, include_md)
        if source.stat().st_size != item["size"] or sha256_file(source) != original_hash:
            raise ValueError("El original cambió durante el procesamiento; no se publican derivados.")
        publish_bundle(stage, source, desired)
        return {"source": str(source), "role": role, "status": "GENERADO", "outputs": [str(output_paths(source)[key]) for key in sorted(desired)], "original_modified": False}
    finally:
        shutil.rmtree(stage, ignore_errors=True)


def process(target: Path, enrichment_path: Path | None = None, forced_failure: bool = False) -> dict[str, Any]:
    target = target.resolve()
    if OUTPUT_80 in target.parts:
        return {"estado": "BLOQUEADO", "modo": PROCESSING_MODE, "motivo": f"{OUTPUT_80} no puede ser fuente.", "acciones_ejecutadas": "Ninguna."}
    if not target.exists() or target.is_symlink() or not (target.is_file() or target.is_dir()):
        raise ValueError("La ruta debe ser un archivo o carpeta ordinaria existente.")
    root, descriptors = prepare_descriptors(target)
    enrichment = load_enrichment(enrichment_path)
    results = []
    for item in descriptors:
        results.append(process_one(item, root, enrichment, forced_failure))
    return {
        "estado": "COMPLETADO", "modo": PROCESSING_MODE, "ruta_analizada": str(target),
        "recursos_detectados": len(descriptors), "resultados": results,
        "originales_modificados": False, "version_skill": SKILL_VERSION,
        "nota": "No se clasificaron, movieron, renombraron, fusionaron ni eliminaron originales.",
    }


def dry_run(target: Path) -> dict[str, Any]:
    target = target.resolve()
    if OUTPUT_80 in target.parts:
        return {"estado": "BLOQUEADO", "modo": PROCESSING_MODE, "motivo": f"{OUTPUT_80} no puede ser fuente."}
    root, descriptors = prepare_descriptors(target)
    return {
        "estado": "APROBADO", "modo": PROCESSING_MODE, "ruta_analizada": str(target),
        "recursos": [
            {"ruta_relativa": item["path"].relative_to(root).as_posix() if root.is_dir() else item["path"].name, "role": item["role"], "basis": item["basis"]}
            for item in descriptors
        ],
        "destino": "DERIVADOS_LATERALES_JUNTO_A_CADA_ORIGINAL",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("ruta", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--ai-enrichment", type=Path)
    parser.add_argument("--force-validation-failure", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    try:
        report = dry_run(args.ruta) if args.dry_run else process(args.ruta, args.ai_enrichment, args.force_validation_failure)
    except Exception as exc:
        report = {"estado": "ERROR", "modo": PROCESSING_MODE, "motivo": f"{type(exc).__name__}: {exc}", "acciones": "No se publicaron derivados incompletos."}
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 2 if report["estado"] == "BLOQUEADO" else 0


if __name__ == "__main__":
    raise SystemExit(main())
